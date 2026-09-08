"""
Divorce Petition Generator — Section 13B Hindu Marriage Act, 1955
Adv. Sachin Madhukar Mahajan

Usage:
  Fill in the CLIENT DATA section below, then run:
      python generate_petition.py

Output:
  A .docx file saved to the project root.
"""

import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

# =============================================================================
# CLIENT DATA — Fill these before running
# =============================================================================

COURT_CITY         = "अमळनेर"                         # e.g. अमळनेर / जळगाव / धुळे
COURT_NAME         = "मे. दिवाणी न्यायाधीश वरिष्ठ स्तर"

HMP_NO             = ""                                 # Case number — leave blank if not yet filed
YEAR               = "२०२५"

# Husband (Applicant No. 1)
H_PREFIX           = "श्री."
H_FULL_NAME        = ""                                 # e.g. नितीन मधुकर महाजन
H_AGE              = ""                                 # e.g. ४५
H_OCCUPATION       = ""                                 # e.g. सेवा (नोकरी) / शेती / व्यवसाय
H_ADDRESS          = ""                                 # e.g. अकुलखेडा, ता. चोपडा, जि. जळगाव

# Wife (Applicant No. 2)
W_PREFIX           = "सौ."
W_MARRIED_NAME     = ""                                 # e.g. संजीवनी नितीन महाजन
W_MAIDEN_NAME      = ""                                 # e.g. संजीवनी राधेश्याम महाजन (before marriage)
W_AGE              = ""                                 # e.g. ३०
W_OCCUPATION       = ""                                 # e.g. सेवा (नोकरी) / गृहिणी
W_GUARDIAN         = ""                                 # Father's name for "Care of / द्वारा"
W_ADDRESS          = ""                                 # e.g. द्वारा राधेश्याम महाजन, अकुलखेडा, ता. चोपडा, जि. जळगाव

# Marriage
MARRIAGE_DATE      = ""                                 # e.g. ०१/०१/२००५
MARRIAGE_PLACE     = ""                                 # e.g. जळगाव

# Children (leave empty list if none)
# Format: [{"name": "...", "age": "...", "custody": "wife/husband"}]
CHILDREN           = []

# Separation
SEPARATION_DATE    = "     /    /        "             # e.g. ११/१२/२०२० or partial "     /    /२०२०"
SEPARATION_YEARS   = ""                                 # e.g. ५ वर्षांपेक्षा जास्त

# Settlement
ALIMONY_AMOUNT_NUM = ""                                 # e.g. १,००,०००
ALIMONY_AMOUNT_WORDS = ""                               # e.g. रुपये एक लाख मात्र
ALIMONY_NIL        = False                              # Set True if no alimony

# Custody (used only if CHILDREN is not empty)
CUSTODY_WITH       = "wife"                             # "wife" or "husband"

# Advocates
ADV_APP1           = "ॲड. सचिन मधुकर महाजन"
ADV_APP2           = ". . . . . . . . . . . . . . . ."

# Output file
OUTPUT_DIR         = os.path.join(os.path.dirname(__file__), "..", "..", "..", "..", "..", "..")
OUTPUT_FILENAME    = f"{H_FULL_NAME.replace(' ', '_')}_{W_MARRIED_NAME.replace(' ', '_')}_Divorce_Petition_13B.docx"

# =============================================================================
# DOCUMENT GENERATION — Do not edit below unless changing format
# =============================================================================

doc = docx.Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1.2)
    section.right_margin  = Inches(1)


def para(text, bold=False, italic=False, underline=False,
         align=WD_ALIGN_PARAGRAPH.LEFT, size=12, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold      = bold
    r.italic    = italic
    r.underline = underline
    r.font.size = Pt(size)
    r.font.name = "Mangal"
    return p


def blank():
    doc.add_paragraph()


# ── Header ───────────────────────────────────────────────────────────────────
para(f"{COURT_CITY} येथील {COURT_NAME} यांचे कोर्टात.",
     bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
para(f"HMP / विवाह अर्ज / नंबर  {HMP_NO}  /{YEAR}",
     bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
blank()

# ── Applicant 1 (Husband) ────────────────────────────────────────────────────
para(f"{H_PREFIX} {H_FULL_NAME}", bold=True)
para(f"वय {H_AGE} वर्षे, धंदा :- {H_OCCUPATION}")
para(f"रा. {H_ADDRESS}.")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r1 = p.add_run("                                               आणि.")
r1.font.name = "Mangal"; r1.font.size = Pt(12)
r2 = p.add_run("                     ….......... अर्जदार क्र. १")
r2.font.name = "Mangal"; r2.font.size = Pt(12)
blank()

# ── Applicant 2 (Wife) ───────────────────────────────────────────────────────
para(f"{W_PREFIX} {W_MARRIED_NAME}", bold=True)
if W_MAIDEN_NAME:
    para(f"उर्फ {W_MAIDEN_NAME}")
para(f"वय {W_AGE} वर्षे, धंदा :- {W_OCCUPATION}")
para(f"रा. {W_ADDRESS}.")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("……….. अर्जदार क्र. २")
r.font.name = "Mangal"; r.font.size = Pt(12)
blank()

para("             //विरुद्ध//", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para("       कोणीही नाही                                                       . . . . . . सामनेवाले")
blank()

# ── Title ────────────────────────────────────────────────────────────────────
para("हिंदू विवाह कायदा १९५५ चे कलम १३ (ब) प्रमाणे",
     bold=True, underline=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
para("घटस्फोट मिळण्यासाठी अर्ज",
     bold=True, underline=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
blank()

para("अर्जदार मे. कोर्टास विनंती पूर्वक कळवितात की,", size=12)
blank()

# ── Body Paragraphs ──────────────────────────────────────────────────────────

# Determine children text
no_children = (CHILDREN is None or len(CHILDREN) == 0)

if no_children:
    children_text = "सदर विवाहापासून उभयतांना कोणतेही अपत्य नाही."
else:
    names = " आणि ".join([f"{c['name']} वय {c['age']} वर्ष" for c in CHILDREN])
    children_text = f"सदर विवाहापासून उभयतांना अपत्य – {names} आहे/आहेत."

body = [
    ("१)", f"अर्जदार क्र. १ {H_PREFIX} {H_FULL_NAME} व अर्जदार क्र. २ {W_PREFIX} {W_MARRIED_NAME} यांचा विवाह हिंदू धर्मशास्त्राप्रमाणे दिनांक {MARRIAGE_DATE} रोजी {MARRIAGE_PLACE} येथे झालेला आहे. {children_text} तसेच आज रोजी अर्जदार क्र. २ ही गर्भवती नाही."),

    ("२)", "लग्नानंतर अर्जदार क्र. २ ही अर्जदार क्र. १ यांच्याकडे नांदणेस गेली. त्यानंतर उभयतांनी काही काळ संसार केलेला आहे. परंतु दोघांच्या भिन्न स्वभावामुळे व एकमेकांशी पटत नसल्यामुळे व सतत किरकोळ घरगुती कारणांवरून वाद होऊ लागल्याने दोघांना एकत्रित राहून संसार करणे दुरापास्त झाले."),

    ("३)", f"अर्जदार क्र. २ ही दिनांक {SEPARATION_DATE} पासून सुमारे {SEPARATION_YEARS} कालावधीपासून माहेरी राहत आहे. या काळात दोन्ही अर्जदारांच्या कुटुंबातील लोकांनी व नातेवाईकांनी तसेच वकीलांनी अर्जदारांनी संसार करावा म्हणून अनेक प्रयत्न केले परंतु काही उपयोग झाला नाही. भविष्यातील अघटित घटना घडू नये तसेच अर्जदारांनी आपले उर्वरित पुढील आयुष्य हे चांगले जावे म्हणून सदर घटस्फोटाचा निर्णय घेण्याचे ठरवलेले आहे. अर्जदारांचा संसार सुखासमाधानाने भविष्यात होऊ शकत नाही, तसेच अर्जदार एकत्र येणार नाहीत, अशी अर्जदारांची खात्री झाल्याने अर्जदारांनी त्यांचे विवाह संबंध संपुष्टात आणण्याचे ठरवले आहे, आणि दोन्ही अर्जदार घटस्फोट घेण्याच्या निर्णयावर ठाम होते व आहेत. म्हणून हा घटस्फोटाचा अर्ज मे. कोर्टात दाखल करणे भाग झाले आहे."),

    ("४)", "दोन्ही अर्जदारांचा झालेला विवाह हा पुनर्स्थापित होणे अशक्य आहे. अर्जदारांनी आपले भविष्यातील प्रश्न व एकमेकांमध्ये असलेले वाद-विवाद व सर्व समस्या या अर्जदारांचे आई-वडील व नातेवाईक यांच्या मध्यस्थीने सोडवलेल्या आहेत. अर्जदारांमध्ये कोणताही वाद किंवा प्रश्न प्रलंबित नाही."),
]

for num, text in body:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r1 = p.add_run(num + "  ")
    r1.bold = True; r1.font.size = Pt(12); r1.font.name = "Mangal"
    r2 = p.add_run(text)
    r2.font.size = Pt(12); r2.font.name = "Mangal"
    blank()

# ── Settlement Terms ──────────────────────────────────────────────────────────
p = doc.add_paragraph()
r = p.add_run("५)  उभयतांमध्ये खालील अटी व शर्तींच्या आधारे आपसात घटस्फोट घेण्याचे ठरले आहे ते खालीलप्रमाणे:-")
r.bold = True; r.font.size = Pt(12); r.font.name = "Mangal"

# Build settlement clauses dynamically
settlement = []

# Custody clause
if no_children:
    settlement.append("अ)  दोन्ही अर्जदारांना सदर विवाहापासून कोणतेही अपत्य नसल्यामुळे मुलांच्या ताब्याचा कोणताही प्रश्न उद्भवत नाही.")
else:
    custody_person = W_MARRIED_NAME if CUSTODY_WITH == "wife" else H_FULL_NAME
    child_names = " आणि ".join([c["name"] for c in CHILDREN])
    settlement.append(f"अ)  दोन्ही अर्जदारांची अपत्य {child_names} {custody_person} यांच्याकडे राहावयाची आहेत. त्यांचे पालन पोषण व संगोपनाची संपूर्ण जबाबदारी {custody_person} यांची राहील.")

# Alimony clause
if ALIMONY_NIL or not ALIMONY_AMOUNT_NUM:
    settlement.append(f"ब)  अर्जदार क्रमांक २ {W_PREFIX} {W_MARRIED_NAME} यांनी खावटी / पोटगी मागण्याचा संपूर्ण हक्क कायमस्वरूपी विनामोबदला या घटस्फोटापासून स्वखुशीने सोडून दिलेला आहे. भविष्यात अर्जदार क्रमांक २ ही अर्जदार क्रमांक १ किंवा त्यांच्या कुटुंबीयांविरुद्ध कोणतीही खावटी मागणार नाही.")
else:
    settlement.append(f"ब)  अर्जदार क्रमांक १ {H_PREFIX} {H_FULL_NAME} हे अर्जदार क्रमांक २ {W_PREFIX} {W_MARRIED_NAME} हिला मागील, पुढील व भविष्यातील खावटी / पोटगी म्हणून एकरकमी रक्कम रुपये {ALIMONY_AMOUNT_NUM}/- ({ALIMONY_AMOUNT_WORDS}) देण्याचे ठरले आहे. सदर रक्कम मिळाल्यानंतर अर्जदार क्रमांक २ हिला कोणतीही तक्रार राहणार नाही. भविष्यात अर्जदार क्रमांक २ ही अर्जदार क्रमांक १ किंवा त्यांच्या कुटुंबीयांविरुद्ध कोणतीही खावटी मागणार नाही किंवा कोणतेही अर्ज-फाटे-केसेस करणार नाही.")

settlement += [
    "क)  दोन्ही अर्जदार यांची एकमेकांकडे सोने, नाणे, दागदागिने, कपडे-लत्ते देणे-घेणे काहीही बाकी नाही.",
    "ड)  दोन्ही अर्जदारांनी एकमेकांच्या स्थावर व जंगम मिळकतीवर आणि वडीलोपार्जित मिळकतीवर कोणताही हक्क व अधिकार भविष्यात दाखवायचा नाही.",
    "इ)  दोन्ही अर्जदारांनी एकमेकांच्या विरुद्ध आणि एकमेकांच्या कुटुंबीयांच्या विरुद्ध कोणत्याही फौजदारी स्वरूपाच्या केसेस व खटले दाखल केलेले नाहीत आणि भविष्यात देखील करावयाच्या नाहीत.",
    "फ)  सदरच्या अर्जाचे कामकाज चालू असताना आणि भविष्यातही दोन्ही अर्जदारांनी एकमेकांशी प्रत्यक्ष संपर्क साधायचा नाही. संपर्क साधण्यासाठी नातेवाईक आणि वकिलांच्या मदतीने आवश्यक बाबींसाठी संपर्क साधला जाईल.",
    "ग)  दोन्ही अर्जदार यांनी लग्नाच्या संसारादरम्यानचे किंवा इतर कोणत्याही प्रकारचे फोटो व व्हिडिओ कोणत्याही प्रकारे प्रसारित करावयाचे नाहीत. एकमेकांचे नातेवाईक व कुटुंबाविषयी बदनामीकारक किंवा अपशब्द वापर व्हायचे नाहीत.",
]

for clause in settlement:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.3)
    r = p.add_run(clause)
    r.font.size = Pt(12); r.font.name = "Mangal"

blank()

# ── High Court Guidelines ─────────────────────────────────────────────────────
p = doc.add_paragraph()
r = p.add_run("६)  मा. उच्च न्यायालयाच्या आदेशाप्रमाणे आवश्यक माहिती –")
r.bold = True; r.font.size = Pt(12); r.font.name = "Mangal"

hc_info = [
    f"•  दोन्ही अर्जदार भारतीय असून हिंदूधर्मीय आहेत. दोन्ही अर्जदारांनी आपल्या जाती-धर्मात कोणताही बदल केलेला नाही.",
    f"•  अर्जदार क्र. १ यांच्या नावात कोणताही बदल नाही. अर्जदार क्र. २ यांचे लग्नापूर्वीचे नाव {W_MAIDEN_NAME} असून लग्नानंतरचे नाव {W_PREFIX} {W_MARRIED_NAME} असे आहे.",
    f"•  अर्जदारांना सदर लग्न संबंधातून {'कोणतेही अपत्य नाही' if no_children else 'अपत्य आहे/आहेत'}.",
    f"•  दोन्ही अर्जदार दिनांक {SEPARATION_DATE} पासून म्हणजे {SEPARATION_YEARS} कालावधी एकमेकांपासून विभक्त राहत आहेत.",
    "•  अर्जदारांनी सदर अर्जाव्यतिरिक्त घटस्फोट मिळण्यासाठी इतर कोणत्याही न्यायालयात अर्ज दाखल केलेला नाही.",
    "•  सदरचा अर्ज दोन्ही अर्जदारांनी स्वखुशीने दाखल केलेला असून त्यांच्यावर कोणताही दबाव किंवा दडपण नाही.",
]

for info in hc_info:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.3)
    r = p.add_run(info)
    r.font.size = Pt(12); r.font.name = "Mangal"

blank()

# ── Cause of Action ───────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r1 = p.add_run("७)  अर्जास कारण – ")
r1.bold = True; r1.font.size = Pt(12); r1.font.name = "Mangal"
r2 = p.add_run(f"अर्जदार क्रमांक १ व २ यांच्यात वैचारिक मतभेद आहेत म्हणून दिनांक {SEPARATION_DATE} पासून अर्जदार विभक्त असून, एकमेकांत तडजोड करून संसार करण्याचे अनेक प्रयत्न अयशस्वी झाले. तसेच अर्जदारांचा पुन्हा एकत्र येऊन संसार होणे अशक्य असल्याने सदर अर्जास कारण घडले व सतत घडत आहे.")
r2.font.size = Pt(12); r2.font.name = "Mangal"
blank()

# ── Jurisdiction ──────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r1 = p.add_run("८)  न्यायाक्षेत्र – ")
r1.bold = True; r1.font.size = Pt(12); r1.font.name = "Mangal"
r2 = p.add_run(f"अर्जदारांचा विवाह {MARRIAGE_PLACE} येथे झालेला असून अर्जदार क्र. २ सध्या {W_ADDRESS} येथे मे. कोर्टाच्या स्थळसीमेत रहिवास करीत असल्याने मे. कोर्टास सदर अर्ज चालवण्याचा पूर्ण अधिकार आहे.")
r2.font.size = Pt(12); r2.font.name = "Mangal"
blank()

# ── Court Fees ────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r1 = p.add_run("९)  कोर्ट फी – ")
r1.bold = True; r1.font.size = Pt(12); r1.font.name = "Mangal"
r2 = p.add_run("अर्जदारांनी सदर अर्जास योग्य तो कोर्ट फी स्टॅम्प लावलेला आहे.")
r2.font.size = Pt(12); r2.font.name = "Mangal"

blank(); blank()

# ── Prayer ────────────────────────────────────────────────────────────────────
para("तरी अर्जदारांची विनंती की,", bold=True)
blank()

prayers = [
    f"अ)  अर्जदार यांचा अर्ज मंजूर करण्यात येऊन दिनांक {MARRIAGE_DATE} रोजी अर्जदारांचा झालेला विवाह विच्छेदन करण्यात यावा. आणि अर्जदारांचा घटस्फोट झाला आहे असे जाहीर करण्यात यावे ही विनंती.",
    "ब)  अर्जदारांच्या हिताचे व न्यायाचे अन्य हुकुम व्हावेत ही विनंती.",
]
for pr in prayers:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(pr)
    r.font.size = Pt(12); r.font.name = "Mangal"

blank(); blank()

# ── Signature Block ───────────────────────────────────────────────────────────
para(f"दिनांक       /     /{YEAR}                                          स्थळ :- {COURT_CITY}")
blank()

for label in [f"अर्जदार नं. १ ({H_PREFIX} {H_FULL_NAME})", f"अर्जदार नं. २ ({W_PREFIX} {W_MARRIED_NAME})"]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(". . . . . . . . . . . . . . . . . . . . . . .")
    r.font.name = "Mangal"; r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(label)
    r.font.name = "Mangal"; r.font.size = Pt(12)
    blank()

# ── Affidavit ─────────────────────────────────────────────────────────────────
doc.add_page_break()

para("प्रतिज्ञालेख", bold=True, underline=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
blank()

aff_text = (
    f"आम्ही १) {H_PREFIX} {H_FULL_NAME} वय {H_AGE} वर्षे, धंदा :- {H_OCCUPATION}, रा. {H_ADDRESS}. "
    f"आणि २) {W_PREFIX} {W_MARRIED_NAME} वय {W_AGE} वर्षे, धंदा :- {W_OCCUPATION}, रा. {W_ADDRESS} – "
    "सत्य प्रतिज्ञेवर कथन करतो की वरील मजकूर आमच्या माहिती व समजुतीप्रमाणे खरा व बरोबर असून "
    "त्यासोबत जोडलेली सहाय्यक कागदपत्रे ही खरी व बरोबर आहेत."
)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r = p.add_run(aff_text)
r.font.size = Pt(12); r.font.name = "Mangal"

blank(); blank()

para(f"दिनांक       /     /{YEAR}                                          स्थळ :- {COURT_CITY}")
blank()

for label in ["अर्जदार नं. १", "अर्जदार नं. २"]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(". . . . . . . . . . . . . . . . . . . . . . .")
    r.font.name = "Mangal"; r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(label)
    r.font.name = "Mangal"; r.font.size = Pt(12)
    blank()

blank()

para(f"अर्जदार क्र. १ तर्फे विधीज्ञ. {ADV_APP1}")
para(f"अर्जदार क्र. २ तर्फे विधीज्ञ. {ADV_APP2}")

# ── Save ──────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.abspath(OUTPUT_DIR), OUTPUT_FILENAME)
doc.save(out_path)
print(f"✓ Petition saved: {out_path}")
